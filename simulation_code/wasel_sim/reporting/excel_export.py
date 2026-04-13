from __future__ import annotations

from pathlib import Path
from typing import Any

import pandas as pd
import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
SECTION_FILL = PatternFill("solid", fgColor="D9EAF7")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")
BODY_FONT = Font(name="Arial", color="000000")
INPUT_FONT = Font(name="Arial", color="0000FF")


def _flatten(prefix: str, value: Any, rows: list[tuple[str, Any]]) -> None:
    if isinstance(value, dict):
        for key, child in value.items():
            new_prefix = f"{prefix}.{key}" if prefix else str(key)
            _flatten(new_prefix, child, rows)
    elif isinstance(value, list):
        rows.append((prefix, ", ".join(str(item) for item in value)))
    else:
        rows.append((prefix, value))


def _apply_table_style(ws) -> None:
    ws.freeze_panes = "A2"
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = BODY_FONT
            if isinstance(cell.value, (int, float)):
                cell.number_format = '#,##0.00;(#,##0.00);-'
    for col_idx, column in enumerate(ws.columns, start=1):
        max_len = 0
        for cell in column:
            value = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, len(value))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 12), 40)
    ws.auto_filter.ref = ws.dimensions


def _write_dataframe(ws, df: pd.DataFrame) -> None:
    ws.append(list(df.columns))
    for row in df.itertuples(index=False, name=None):
        ws.append(list(row))
    _apply_table_style(ws)


def build_excel_workbook(
    out_path: Path,
    assumptions_path: Path,
    summary_df: pd.DataFrame,
    sim_df: pd.DataFrame,
    sensitivity_detail_df: pd.DataFrame,
    sensitivity_rank_df: pd.DataFrame,
    mc_df: pd.DataFrame,
    mc_stats: dict[str, Any],
) -> Path:
    wb = Workbook()
    summary_ws = wb.active
    summary_ws.title = "Summary"

    summary_ws["A1"] = "Wasel Simulation Summary"
    summary_ws["A1"].font = Font(name="Arial", bold=True, size=14)
    summary_ws["A3"] = "Model changes"
    summary_ws["A3"].fill = SECTION_FILL
    summary_ws["A3"].font = Font(name="Arial", bold=True)
    changes = [
        "Revenue comes only from driver-side monetization; rider subscriptions are excluded.",
        "Launch demand is geographically constrained early and expands over time.",
        "Competitor pressure now raises rider CAC, slows organic growth, and trims share capture early.",
        "Rider demand supports utilization with a tolerance buffer and explicit constrained-demand tracking.",
        "LTRC license is expensed monthly in P&L and paid as a cash outflow on the renewal month.",
        "Operational break-even is defined as contribution margin turning positive.",
    ]
    for idx, text in enumerate(changes, start=4):
        summary_ws[f"A{idx}"] = text
        summary_ws[f"A{idx}"].font = BODY_FONT

    metric_row = 8
    summary_ws[f"A{metric_row}"] = "Deterministic headline metrics"
    summary_ws[f"A{metric_row}"].fill = SECTION_FILL
    summary_ws[f"A{metric_row}"].font = Font(name="Arial", bold=True)
    for offset, (col, value) in enumerate(summary_df.iloc[0].items(), start=1):
        row = metric_row + offset
        summary_ws[f"A{row}"] = col
        summary_ws[f"B{row}"] = value
        summary_ws[f"A{row}"].font = BODY_FONT
        summary_ws[f"B{row}"].font = BODY_FONT
        if isinstance(value, (int, float)):
            summary_ws[f"B{row}"].number_format = '#,##0.00;(#,##0.00);-'

    sens_row = metric_row + len(summary_df.columns) + 3
    summary_ws[f"A{sens_row}"] = "Top sensitivity drivers"
    summary_ws[f"A{sens_row}"].fill = SECTION_FILL
    summary_ws[f"A{sens_row}"].font = Font(name="Arial", bold=True)
    top_sens = sensitivity_rank_df.head(10)[["variable", "swing"]]
    for idx, row in enumerate(top_sens.itertuples(index=False), start=1):
        summary_ws[f"A{sens_row + idx}"] = row.variable
        summary_ws[f"B{sens_row + idx}"] = row.swing
        summary_ws[f"B{sens_row + idx}"].number_format = '#,##0.00;(#,##0.00);-'

    mc_row = sens_row + len(top_sens) + 3
    summary_ws[f"A{mc_row}"] = "Monte Carlo summary"
    summary_ws[f"A{mc_row}"].fill = SECTION_FILL
    summary_ws[f"A{mc_row}"].font = Font(name="Arial", bold=True)
    current_row = mc_row + 1
    for key, value in mc_stats.items():
        if isinstance(value, dict):
            summary_ws[f"A{current_row}"] = key
            summary_ws[f"A{current_row}"].font = Font(name="Arial", bold=True)
            current_row += 1
            for stat_key, stat_value in value.items():
                summary_ws[f"A{current_row}"] = stat_key
                summary_ws[f"B{current_row}"] = stat_value
                summary_ws[f"B{current_row}"].number_format = '#,##0.00;(#,##0.00);-'
                current_row += 1
        else:
            summary_ws[f"A{current_row}"] = key
            summary_ws[f"B{current_row}"] = value
            if isinstance(value, (int, float)):
                summary_ws[f"B{current_row}"].number_format = '0.0%' if 0 <= value <= 1 else '#,##0.00;(#,##0.00);-'
            current_row += 1

    for col in ("A", "B", "C"):
        summary_ws.column_dimensions[col].width = 32

    assumptions_ws = wb.create_sheet("Assumptions")
    with open(assumptions_path, "r", encoding="utf-8") as handle:
        assumptions = yaml.safe_load(handle)
    rows: list[tuple[str, Any]] = []
    _flatten("", assumptions, rows)
    assumptions_df = pd.DataFrame(rows, columns=["assumption", "value"])
    _write_dataframe(assumptions_ws, assumptions_df)
    for cell in assumptions_ws["B"][1:]:
        cell.font = INPUT_FONT

    monthly_ws = wb.create_sheet("Monthly Simulation")
    _write_dataframe(monthly_ws, sim_df)

    sensitivity_rank_ws = wb.create_sheet("Sensitivity Rank")
    _write_dataframe(sensitivity_rank_ws, sensitivity_rank_df)

    sensitivity_detail_ws = wb.create_sheet("Sensitivity Detail")
    _write_dataframe(sensitivity_detail_ws, sensitivity_detail_df)

    mc_summary_rows = []
    for key, value in mc_stats.items():
        if isinstance(value, dict):
            for stat_key, stat_value in value.items():
                mc_summary_rows.append({"metric": key, "stat": stat_key, "value": stat_value})
        else:
            mc_summary_rows.append({"metric": key, "stat": "value", "value": value})
    mc_summary_ws = wb.create_sheet("Monte Carlo Summary")
    _write_dataframe(mc_summary_ws, pd.DataFrame(mc_summary_rows))

    mc_raw_ws = wb.create_sheet("Monte Carlo Raw")
    _write_dataframe(mc_raw_ws, mc_df)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path
